"""Генерация клиентського xlsx-звіту про підрахунок (українською).

Чисті функції: DTO з report_data_service -> bytes. Без БД і мережі.

  * місто    — один аркуш «Звіт»;
  * міжмісто — «Підсумок», «Ходки», «Поїздки» (з цінами тарифу), «Відрізки»;
  * загальний звіт за період — один аркуш «Виручка» (машини × дні).
"""

from __future__ import annotations

import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.report_data_service import (
    REVENUE_DESCRIPTION,
    CityReport,
    IntercityReport,
    OverviewReport,
    ReportHeader,
    fmt_date,
    fmt_datetime,
    fmt_duration,
)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_TITLE_FONT = Font(size=14, bold=True)
_BOLD = Font(bold=True)
_BIG_BOLD = Font(size=12, bold=True)
_UPDATE_FONT = Font(bold=True, color="C62828")
_NOTE_FONT = Font(italic=True, color="555555")
_HEAD_FILL = PatternFill("solid", fgColor="E0F2F1")
_TOTAL_FILL = PatternFill("solid", fgColor="F5F5F5")
_RUN_FILL = PatternFill("solid", fgColor="FAFAFA")
_THIN = Side(style="thin", color="BDBDBD")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WRAP = Alignment(wrap_text=True, vertical="top")
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

MONEY_FMT = '#,##0.00 "₴"'
KM_FMT = "0.00"
INT_FMT = "0"
TIME_FMT = "HH:MM"
DATE_FMT = "DD.MM.YYYY"
DELTA_FMT = "+0;-0;0"

# Ширина колонок «текстового» блоку аркуша-зведення.
_LABEL_WIDTH = 34
_VALUE_WIDTH = 70


def report_filename(header: ReportHeader) -> str:
    """Звіт_2026-10-20_<авто>[_версія_2].xlsx"""
    car = re.sub(r"[^\w-]+", "_", header.car_name).strip("_")[:40] or header.imei
    name = f"Звіт_{header.day.isoformat()}_{car}"
    if header.is_update:
        name += f"_версія_{header.version}"
    return f"{name}.xlsx"


def overview_filename(report: OverviewReport) -> str:
    """Виручка_2026-10-01_2026-10-29_<акаунт>.xlsx"""
    title = re.sub(r"[^\w-]+", "_", report.title).strip("_")[:40]
    name = f"Виручка_{report.from_day.isoformat()}_{report.to_day.isoformat()}"
    if title:
        name += f"_{title}"
    return f"{name}.xlsx"


# ── загальні блоки ───────────────────────────────────────────────────

def _text_row_height(text: str, chars_per_line: int = 90) -> float:
    lines = max(1, -(-len(text) // chars_per_line))
    return 15 * lines + 3


def _kv(
    ws: Worksheet, row: int, label: str, value, *,
    value_font=None, number_format=None, wrap: bool = True,
) -> int:
    ws.cell(row=row, column=1, value=label).font = _BOLD
    cell = ws.cell(row=row, column=2, value=value)
    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=wrap)
    if value_font is not None:
        cell.font = value_font
    if number_format:
        cell.number_format = number_format
    return row + 1


def _paragraph(ws: Worksheet, row: int, label: str, text: str, *, font=None) -> int:
    ws.cell(row=row, column=1, value=label).font = _BOLD
    ws.cell(row=row, column=1).alignment = _WRAP
    cell = ws.cell(row=row, column=2, value=text)
    cell.alignment = _WRAP
    if font is not None:
        cell.font = font
    ws.row_dimensions[row].height = _text_row_height(text)
    return row + 1


def _write_header(ws: Worksheet, header: ReportHeader) -> int:
    """Шапка звіту; повертає номер наступного вільного рядка."""
    ws.column_dimensions["A"].width = _LABEL_WIDTH
    ws.column_dimensions["B"].width = _VALUE_WIDTH

    title = "Звіт про підрахунок пасажирів"
    if header.is_update:
        title += f" (версія {header.version}, оновлено)"
    ws.cell(row=1, column=1, value=title).font = _TITLE_FONT

    row = 3
    row = _kv(ws, row, "Транспортний засіб", header.car_name)
    row = _kv(ws, row, "IMEI", header.imei)
    if header.company:
        row = _kv(ws, row, "Перевізник", header.company)
    row = _kv(ws, row, "Дата підрахунку", fmt_date(header.day))
    row = _kv(ws, row, "Тип маршруту", header.kind_label)
    row = _kv(
        ws, row, "Версія звіту",
        f"{header.version} (оновлена)" if header.is_update else "1",
        value_font=_UPDATE_FONT if header.is_update else None,
    )
    if header.is_update:
        row = _paragraph(ws, row, "Увага", header.update_note, font=_UPDATE_FONT)
    row = _kv(ws, row, "Звіт сформовано", fmt_datetime(header.generated_at))
    row += 1
    row = _paragraph(ws, row, "Що підраховано", header.counted_description)
    return row + 1


def _write_table(
    ws: Worksheet,
    start_row: int,
    columns: list[tuple[str, float, str | None]],
    rows: list[list],
    *,
    total: list | None = None,
    row_fills: list | None = None,
    row_fonts: list | None = None,
) -> int:
    """Таблиця з шапкою. columns: [(назва, ширина, number_format)]."""
    for idx, (title, width, _fmt) in enumerate(columns, start=1):
        cell = ws.cell(row=start_row, column=idx, value=title)
        cell.font = _BOLD
        cell.fill = _HEAD_FILL
        cell.border = _BORDER
        cell.alignment = _CENTER
        letter = get_column_letter(idx)
        current = ws.column_dimensions[letter].width or 0
        ws.column_dimensions[letter].width = max(current, width)

    row = start_row + 1
    for i, values in enumerate(rows):
        fill = row_fills[i] if row_fills else None
        font = row_fonts[i] if row_fonts else None
        for idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=idx, value=value)
            cell.border = _BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(value, str))
            fmt = columns[idx - 1][2]
            if fmt and not isinstance(value, str):
                cell.number_format = fmt
            if fill is not None:
                cell.fill = fill
            if font is not None:
                cell.font = font
        row += 1

    if total is not None:
        for idx, value in enumerate(total, start=1):
            cell = ws.cell(row=row, column=idx, value=value)
            cell.font = _BOLD
            cell.fill = _TOTAL_FILL
            cell.border = _BORDER
            fmt = columns[idx - 1][2]
            if fmt and not isinstance(value, str):
                cell.number_format = fmt
        row += 1
    return row


def _to_bytes(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _new_workbook(header: ReportHeader) -> Workbook:
    wb = Workbook()
    wb.properties.title = f"Звіт про підрахунок пасажирів {fmt_date(header.day)}"
    wb.properties.creator = "Підрахунок пасажиропотоку"
    return wb


# ── місто ────────────────────────────────────────────────────────────

def build_city_xlsx(report: CityReport) -> bytes:
    wb = _new_workbook(report.header)
    ws = wb.active
    ws.title = "Звіт"

    row = _write_header(ws, report.header)
    row = _kv(ws, row, "Пасажирів за день", report.passengers,
              value_font=_BIG_BOLD, number_format=INT_FMT)
    row = _kv(ws, row, "Пробіг, км", report.distance_km, number_format=KM_FMT)
    if report.window_start and report.window_end:
        period = (
            f"{fmt_datetime(report.window_start)} — "
            f"{report.window_end.strftime('%H:%M')}"
        )
        row = _kv(ws, row, "Період підрахунку", period)
    return _to_bytes(wb)


# ── міжмісто ─────────────────────────────────────────────────────────

def _intercity_summary(ws: Worksheet, report: IntercityReport) -> None:
    row = _write_header(ws, report.header)
    row = _paragraph(ws, row, "Як розраховано виручку", REVENUE_DESCRIPTION, font=_NOTE_FONT)
    row += 1
    row = _kv(ws, row, "Ходок за день", report.run_count, number_format=INT_FMT)
    row = _kv(ws, row, "Час у роботі", fmt_duration(report.work_sec))
    row = _kv(ws, row, "Пробіг, км", report.distance_km, number_format=KM_FMT)
    row = _kv(ws, row, "Пасажирів (посадок) за день", report.passengers,
              value_font=_BIG_BOLD, number_format=INT_FMT)
    row = _kv(ws, row, "Розрахована виручка", report.revenue,
              value_font=_BIG_BOLD, number_format=MONEY_FMT)
    if report.unpriced_pax:
        row = _paragraph(
            ws, row, "Без тарифу",
            f"Для {report.unpriced_pax} пас. не знайдено ціну проїзду в тарифі — "
            "їх не враховано у виручці.",
            font=_UPDATE_FONT,
        )
    if report.tariffs:
        row += 1
        ws.cell(row=row, column=1, value="Застосовані тарифи").font = _BIG_BOLD
        row += 1
        _write_table(
            ws, row,
            [("Маршрут", _LABEL_WIDTH, None), ("Тариф", _VALUE_WIDTH, None)],
            [[t.route_name, t.tariff_name or "тариф не знайдено"] for t in report.tariffs],
        )


def _intercity_runs(ws: Worksheet, report: IntercityReport) -> None:
    columns = [
        ("№", 5, INT_FMT),
        ("Маршрут", 36, None),
        ("Ходка", 8, INT_FMT),
        ("Початок", 10, TIME_FMT),
        ("Кінець", 10, TIME_FMT),
        ("Тривалість", 14, None),
        ("Км по маршруту", 11, KM_FMT),
        ("Посадок", 10, INT_FMT),
        ("Виручка", 14, MONEY_FMT),
        ("Тариф", 30, None),
    ]
    rows = [
        [
            r.number, r.route_name, r.occurrence, r.start, r.end,
            fmt_duration(r.duration_sec), r.distance_km, r.boardings,
            r.revenue if r.revenue is not None else "—",
            r.tariff_name or "тариф не знайдено",
        ]
        for r in report.runs
    ]
    total = [
        "", "Разом", "", "", "", fmt_duration(report.work_sec),
        round(sum(r.distance_km for r in report.runs), 2),
        sum(r.boardings for r in report.runs), report.revenue, "",
    ]
    _write_table(ws, 1, columns, rows, total=total)
    ws.freeze_panes = "A2"


def _intercity_trips(ws: Worksheet, report: IntercityReport) -> None:
    columns = [
        ("№ ходки", 8, INT_FMT),
        ("Маршрут", 34, None),
        ("Звідки (посадка)", 32, None),
        ("Куди (висадка)", 32, None),
        ("Пасажирів", 11, INT_FMT),
        ("Ціна", 12, MONEY_FMT),
        ("Сума", 14, MONEY_FMT),
    ]
    rows = [
        [
            t.run_number, t.route_name, t.from_name, t.to_name, t.pax,
            t.price if t.price is not None else "тариф не знайдено",
            t.subtotal if t.subtotal is not None else "—",
        ]
        for t in report.trips
    ]
    total = ["", "Разом", "", "", sum(t.pax for t in report.trips), "", report.revenue]
    _write_table(ws, 1, columns, rows, total=total)
    ws.freeze_panes = "A2"


def _intercity_segments(ws: Worksheet, report: IntercityReport) -> None:
    columns = [
        ("№ ходки", 8, INT_FMT),
        ("Маршрут", 34, None),
        ("Відправлення", 13, TIME_FMT),
        ("Прибуття", 11, TIME_FMT),
        ("Від зупинки", 32, None),
        ("До зупинки", 32, None),
        ("Відстань, км", 12, KM_FMT),
        ("У салоні", 10, INT_FMT),
        ("Зміна", 9, DELTA_FMT),
    ]
    rows = []
    fills = []
    for s in report.segments:
        rows.append([
            s.run_number, s.route_name, s.departure, s.arrival,
            s.from_name, s.to_name, s.distance_km,
            s.count if s.count is not None else "—", s.delta,
        ])
        fills.append(_RUN_FILL if s.run_number % 2 == 0 else None)
    _write_table(ws, 1, columns, rows, row_fills=fills)
    ws.freeze_panes = "A2"


def build_intercity_xlsx(report: IntercityReport) -> bytes:
    wb = _new_workbook(report.header)
    summary = wb.active
    summary.title = "Підсумок"
    _intercity_summary(summary, report)
    _intercity_runs(wb.create_sheet("Ходки"), report)
    _intercity_trips(wb.create_sheet("Поїздки"), report)
    _intercity_segments(wb.create_sheet("Відрізки"), report)
    return _to_bytes(wb)


def build_xlsx(report: CityReport | IntercityReport) -> bytes:
    if isinstance(report, IntercityReport):
        return build_intercity_xlsx(report)
    return build_city_xlsx(report)


# ── загальний звіт за період ─────────────────────────────────────────

def _money_or_dash(value):
    return value if value is not None else "—"


def build_overview_xlsx(report: OverviewReport) -> bytes:
    wb = Workbook()
    wb.properties.title = (
        f"Виручка {fmt_date(report.from_day)} — {fmt_date(report.to_day)}"
    )
    wb.properties.creator = "Підрахунок пасажиропотоку"
    ws = wb.active
    ws.title = "Виручка"

    ws.cell(row=1, column=1, value="Звіт про пасажирів і виручку").font = _TITLE_FONT
    row = 3
    if report.title:
        row = _kv(ws, row, "Акаунт", report.title, wrap=False)
    row = _kv(ws, row, "Період",
              f"{fmt_date(report.from_day)} — {fmt_date(report.to_day)}", wrap=False)
    row = _kv(ws, row, "Авто з даними", len(report.cars), number_format=INT_FMT, wrap=False)
    row = _kv(ws, row, "Звіт сформовано", fmt_datetime(report.generated_at), wrap=False)
    row += 1
    row = _kv(ws, row, "Пасажирів за період", report.passengers,
              value_font=_BIG_BOLD, wrap=False)
    row = _kv(ws, row, "Виручка за період", _money_or_dash(report.revenue),
              value_font=_BIG_BOLD, number_format=MONEY_FMT, wrap=False)
    row += 1

    if report.cars:
        columns = [
            ("Авто", _LABEL_WIDTH, None),
            ("IMEI", 18, None),
            ("Дата", 14, DATE_FMT),
            ("Кілометраж, км", 14, KM_FMT),
            ("Пасажири", 11, None),
            ("Виручка", 15, MONEY_FMT),
            ("Тип", 11, None),
        ]
        rows, fills, fonts = [], [], []
        for car in report.cars:
            for d in car.days:
                rows.append([
                    car.name, car.imei, d.day, d.distance_km, d.passengers,
                    _money_or_dash(d.revenue), d.kind_label,
                ])
                fills.append(None)
                fonts.append(None)
            rows.append([
                car.name, car.imei, "Разом по авто", "", car.passengers,
                _money_or_dash(car.revenue), "",
            ])
            fills.append(_TOTAL_FILL)
            fonts.append(_BOLD)
        total = [
            "Разом", f"{len(report.cars)} авто", "", "",
            report.passengers, _money_or_dash(report.revenue), "",
        ]
        table_row = row
        row = _write_table(ws, table_row, columns, rows,
                           total=total, row_fills=fills, row_fonts=fonts)
        ws.freeze_panes = f"A{table_row + 1}"
    else:
        ws.column_dimensions["A"].width = _LABEL_WIDTH
        ws.cell(row=row, column=1, value="Немає даних за вибраний період").font = _NOTE_FONT
        row += 1

    if report.missing:
        row += 1
        ws.cell(row=row, column=1, value="Немає даних по авто").font = _BIG_BOLD
        row += 1
        for car in report.missing:
            ws.cell(row=row, column=1, value=car.name)
            ws.cell(row=row, column=2, value=car.imei)
            row += 1

    return _to_bytes(wb)
